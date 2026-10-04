"""Unit tests for the AFA job scraper.

Run from the ``cloud-function`` directory::

    python -m unittest -v

These tests never touch the network — HTTP is mocked at ``job_scraper.requests``.
"""

import io
import os
import unittest
from datetime import datetime, timezone
from unittest.mock import patch

import openpyxl

import excel_generator
import job_scraper
import main


class _FakeResponse:
    """Minimal stand-in for ``requests.Response``."""

    def __init__(self, status_code: int, payload=None, text: str = ""):
        self.status_code = status_code
        self._payload = payload
        self.text = text

    def json(self):
        if self._payload is None:
            raise ValueError("no JSON body")
        return self._payload


class LookbackTests(unittest.TestCase):
    def test_first_run_uses_initial_window(self):
        now = datetime(2026, 9, 6, 12, tzinfo=timezone.utc)
        self.assertEqual(main._lookback_days(None, now), 45)

    def test_partial_days_are_rounded_up_to_avoid_gaps(self):
        now = datetime(2026, 9, 6, 12, tzinfo=timezone.utc)
        self.assertEqual(main._lookback_days("2026-09-05T11:59:59Z", now), 2)

    def test_unsupported_gap_fails_instead_of_advancing_checkpoint(self):
        now = datetime(2026, 9, 6, 12, tzinfo=timezone.utc)
        with self.assertRaisesRegex(ValueError, "at most 100 days"):
            main._lookback_days("2026-05-01T00:00:00Z", now)


class RunEnvTests(unittest.TestCase):
    def test_default_is_local(self):
        with patch.dict(os.environ, {}, clear=True):
            self.assertEqual(main._resolve_run_env(None), "local")

    def test_environment_variable_selects_the_mode(self):
        with patch.dict(os.environ, {"RUN_ENV": "gcp"}, clear=True):
            self.assertEqual(main._resolve_run_env(None), "gcp")

    def test_explicit_argument_wins_and_is_case_insensitive(self):
        self.assertEqual(main._resolve_run_env("GCP"), "gcp")

    def test_unknown_mode_is_rejected(self):
        with self.assertRaises(ValueError):
            main._resolve_run_env("production")


class SalaryFormattingTests(unittest.TestCase):
    def test_range_is_rendered_with_unit(self):
        self.assertEqual(
            job_scraper._format_salary(
                {"verguetungsangabe": "JAHRESGEHALT",
                 "gehaltsspanneVon": 60000.0, "gehaltsspanneBis": 70000.0}
            ),
            "60,000 \u2013 70,000 EUR / year",
        )

    def test_float_noise_is_rounded(self):
        self.assertEqual(
            job_scraper._format_salary(
                {"verguetungsangabe": "JAHRESGEHALT",
                 "gehaltsspanneVon": 70000.0, "gehaltsspanneBis": 74999.97}
            ),
            "70,000 \u2013 75,000 EUR / year",
        )

    def test_hourly_rate_keeps_one_decimal(self):
        self.assertEqual(
            job_scraper._format_salary(
                {"verguetungsangabe": "STUNDENLOHN",
                 "gehaltsspanneVon": 29.5, "gehaltsspanneBis": 36.4}
            ),
            "29.5 \u2013 36.4 EUR / hour",
        )

    def test_equal_bounds_collapse_to_one_value(self):
        self.assertEqual(
            job_scraper._format_salary(
                {"verguetungsangabe": "JAHRESGEHALT",
                 "gehaltsspanneVon": 69000.0, "gehaltsspanneBis": 69000.0}
            ),
            "69,000 EUR / year",
        )

    def test_unknown_period_omits_the_suffix(self):
        self.assertEqual(
            job_scraper._format_salary({"gehaltsspanneVon": 1000.0}), "from 1,000 EUR"
        )

    def test_missing_figures_become_na(self):
        self.assertEqual(job_scraper._format_salary({"verguetungsangabe": "KEINE_ANGABEN"}), "N/A")
        self.assertEqual(job_scraper._format_salary({}), "N/A")


class HttpBehaviourTests(unittest.TestCase):
    """A 403 must not be retried; transient errors must be."""

    @patch("job_scraper.requests.get")
    def test_403_fails_fast_without_retrying(self, get):
        get.return_value = _FakeResponse(403)
        with self.assertRaises(job_scraper.AuthError):
            job_scraper._get_page("radar", 7, 1)
        self.assertEqual(get.call_count, 1)

    @patch("job_scraper.requests.get")
    def test_400_fails_fast(self, get):
        get.return_value = _FakeResponse(400, text="EINGABEN_UNVOLLSTAENDIG")
        with self.assertRaises(job_scraper.ScrapeError):
            job_scraper._get_page("radar", 7, 1)
        self.assertEqual(get.call_count, 1)

    @patch("job_scraper.time.sleep")
    @patch("job_scraper.requests.get")
    def test_transient_error_is_retried_then_succeeds(self, get, _sleep):
        get.side_effect = [_FakeResponse(503), _FakeResponse(200, {"ergebnisliste": []})]
        self.assertEqual(job_scraper._get_page("radar", 7, 1), {"ergebnisliste": []})
        self.assertEqual(get.call_count, 2)

    @patch("job_scraper.requests.get")
    def test_request_uses_the_v6_endpoint(self, get):
        get.return_value = _FakeResponse(200, {"ergebnisliste": []})
        job_scraper._get_page("radar", 7, 1)
        url = get.call_args.args[0]
        self.assertTrue(url.endswith("/pc/v6/jobs"), url)
        self.assertIn("X-API-Key", get.call_args.kwargs["headers"])


class ScraperTests(unittest.TestCase):
    @patch("job_scraper._get_page")
    def test_v6_results_are_normalized(self, get_page):
        get_page.return_value = {
            "maxErgebnisse": 1,
            "ergebnisliste": [
                {
                    "stellenangebotsTitel": "Radar Engineer",
                    "firma": "Example GmbH",
                    "stellenlokationen": [{"adresse": {"ort": "Ulm"}}],
                    "eintrittszeitraum": {"von": "2026-10-01"},
                    "datumErsteVeroeffentlichung": "2026-09-06",
                    "referenznummer": "ABC-123",
                    "verguetungsangabe": "JAHRESGEHALT",
                    "gehaltsspanneVon": 60000.0,
                    "gehaltsspanneBis": 70000.0,
                }
            ],
        }

        jobs = job_scraper.search_jobs("radar", 1)

        self.assertEqual(len(jobs), 1)
        job = jobs[0]
        self.assertEqual(job["refnr"], "ABC-123")
        self.assertEqual(job["ort"], "Ulm")
        self.assertEqual(job["arbeitgeber"], "Example GmbH")
        self.assertEqual(job["eintrittsdatum"], "2026-10-01")
        self.assertEqual(job["veroeffentlichungsdatum"], "2026-09-06")
        self.assertEqual(job["gehalt"], "60,000 \u2013 70,000 EUR / year")
        self.assertTrue(job["url"].endswith("/ABC-123"))

    @patch("job_scraper._get_page")
    def test_all_pages_are_fetched(self, get_page):
        get_page.return_value = {"maxErgebnisse": 250, "ergebnisliste": [{"referenznummer": "x"}]}
        job_scraper.search_jobs("radar", 7)
        self.assertEqual(get_page.call_count, 3)  # 250 results / 100 per page

    def test_lookback_bounds_are_enforced(self):
        with self.assertRaises(ValueError):
            job_scraper.search_jobs("radar", days=101)


class WorkbookTests(unittest.TestCase):
    @staticmethod
    def _job(refnr="A-1", **overrides):
        job = {
            "search_term": "radar",
            "titel": "Radar Engineer",
            "arbeitgeber": "Acme",
            "ort": "Ulm",
            "gehalt": "60,000 \u2013 70,000 EUR / year",
            "refnr": refnr,
            "eintrittsdatum": "2026-10-01",
            "veroeffentlichungsdatum": "2026-09-06",
            "url": f"https://example.test/{refnr}",
        }
        job.update(overrides)
        return job

    def test_create_writes_expected_schema(self):
        wb = openpyxl.load_workbook(
            io.BytesIO(excel_generator.create_new_workbook({"Radar": [self._job()]}))
        )
        ws = wb["Radar"]
        self.assertEqual([c.value for c in ws[1]], excel_generator.HEADERS)
        self.assertEqual(
            [c.value for c in ws[2]],
            ["radar", "Radar Engineer", "Acme", "Ulm", "60,000 \u2013 70,000 EUR / year",
             "https://example.test/A-1", "A-1", "2026-10-01", "2026-09-06"],
        )

    def test_filter_covers_data_rows(self):
        wb = openpyxl.load_workbook(
            io.BytesIO(excel_generator.create_new_workbook({"Radar": [self._job()]}))
        )
        # A header-only range would render dropdowns that filter nothing.
        self.assertEqual(wb["Radar"].auto_filter.ref, "A1:I2")

    def test_link_column_is_clickable(self):
        wb = openpyxl.load_workbook(
            io.BytesIO(excel_generator.create_new_workbook({"Radar": [self._job()]}))
        )
        ws = wb["Radar"]
        self.assertEqual(ws["F2"].hyperlink.target, "https://example.test/A-1")

    def test_append_skips_existing_refnrs(self):
        first = excel_generator.create_new_workbook({"Radar": [self._job("A-1")]})
        merged = excel_generator.append_to_workbook(
            first, {"Radar": [self._job("A-1"), self._job("A-2")]}
        )
        wb = openpyxl.load_workbook(io.BytesIO(merged))
        ws = wb["Radar"]
        self.assertEqual(ws.max_row, 3)  # header + A-1 + A-2
        self.assertEqual(ws["G3"].value, "A-2")
        self.assertEqual(wb["Radar"].auto_filter.ref, "A1:I3")

    def test_append_creates_missing_sheet(self):
        first = excel_generator.create_new_workbook({"Radar": [self._job("A-1")]})
        merged = excel_generator.append_to_workbook(first, {"Perception": [self._job("B-1")]})
        wb = openpyxl.load_workbook(io.BytesIO(merged))
        self.assertIn("Perception", wb.sheetnames)
        self.assertEqual(wb["Perception"]["G2"].value, "B-1")

    def test_sheet_names_are_sanitized(self):
        wb = openpyxl.load_workbook(
            io.BytesIO(excel_generator.create_new_workbook({"A/B:C*": []}))
        )
        self.assertEqual(wb.sheetnames, ["A-B-C-"])


if __name__ == "__main__":
    unittest.main()
